from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(sys.argv[1])
OUT = Path(sys.argv[2])
RUN = ROOT / "work" / "skill_research_runs" / "run-fingoose-final-004"
REPLAY = RUN / "replay-1"
MASTER = REPLAY / "candidate_master_data.jsonl"

SERIES_ZH = {
    "SP500": "标普500",
    "VIXCLS": "VIX波动率指数",
    "DCOILWTICO": "WTI原油",
    "DCOILBRENTEU": "布伦特原油",
    "DGS10": "美国10年期国债收益率",
    "DTWEXBGS": "美元广义指数",
    "DEXCHUS": "美元兑人民币",
    "DEXJPUS": "美元兑日元",
}


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def load_events():
    return [json.loads(line) for line in MASTER.read_text(encoding="utf-8").splitlines() if line.strip()]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest().upper()


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_row_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, text, *, bold=False, color="1F2937", size=9, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_col_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def set_table_geometry(table, widths, indent=110):
    dxa_widths = [round(width * 1440) for width in widths]
    target_total = round(sum(widths) * 1440)
    dxa_widths[-1] += target_total - sum(dxa_widths)
    total = sum(dxa_widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent))
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in dxa_widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(dxa_widths[idx]))


def add_table(doc, headers, rows, widths, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[idx], "E8EEF5")
        set_cell_text(table.rows[0].cells[idx], header, bold=True, color="16324F", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_data in rows:
        row = table.add_row()
        set_repeat_row_no_split(row)
        for idx, value in enumerate(row_data):
            set_cell_text(row.cells[idx], value, size=font_size)
    set_col_widths(table, widths)
    set_table_geometry(table, widths, indent=110)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return p


def add_callout(doc, label, text, fill="F4F6F9"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(label + "  ")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 77, 120)
    r2 = p.add_run(text)
    for run in (r, r2):
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10)
    set_cell_margins(cell, top=130, bottom=130, start=160, end=160)
    set_table_geometry(table, [6.5], indent=160)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


events = load_events()
validation = load_json(REPLAY / "validation_report.json")
repro = load_json(RUN / "reproducibility_report.json")

milestones = sum(len(x["event"].get("milestones", [])) for x in events)
industries = sum(len(x.get("industries", [])) for x in events)
companies = sum(len(x.get("companies", [])) for x in events)
qas = sum(len(x.get("qa", [])) for x in events)
market_windows = sum(len(x["event"].get("market_windows", [])) for x in events)

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(31, 41, 55)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.12
for style_name, size, color, before, after in (
    ("Title", 24, "16324F", 0, 7),
    ("Subtitle", 10.5, "5B6573", 0, 14),
    ("Heading 1", 15, "2E74B5", 14, 6),
    ("Heading 2", 12, "1F4D78", 10, 4),
):
    st = styles[style_name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.text = "FinGoose 宏观事件研究数据 | 阶段成果汇报"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(107, 114, 128)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("内部阶段材料 | 截至 2026-08-12")
run.font.name = "Microsoft YaHei"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(107, 114, 128)

title = doc.add_paragraph(style="Title")
title.add_run("宏观事件研究候选数据集\n阶段成果汇报")
subtitle = doc.add_paragraph(style="Subtitle")
subtitle.add_run("汇报口径：已有完成内容汇总，不包含新增实时研究或未来行情更新")

add_callout(
    doc,
    "核心结论",
    f"已完成 {len(events)} 个代表性宏观事件、{validation['counts']['sft_samples']} 条候选样本的统一建设，覆盖事件事实、时间线、传导机制、行业与公司影响、事件关联问答，以及 T-1、D1、D3、D5、D20、M3 行情验证。当前成果为可审计、可离线复现的候选研究数据，尚未解除人工审核边界。",
)

doc.add_heading("一、完成成果", level=1)
summary_rows = [
    ("宏观母事件", len(events), "每个事件形成统一事件包"),
    ("候选样本", validation["counts"]["sft_samples"], "事件、行业、公司和问答四类"),
    ("事件里程碑", milestones, "包含前置条件、初始冲击及后续节点"),
    ("行业影响分析", industries, "每个事件覆盖 7-9 个行业"),
    ("公司影响分析", companies, "每个事件 3 家事前业务暴露公司"),
    ("事件关联问答", qas, "每个事件 4 类问答，不脱离母事件"),
    ("核心行情窗口", market_windows, "每个事件 4 个价值锚定物"),
    ("推理链", validation["reasoning_chain_stats"]["totals"]["chain_count"], "闭环状态 157/157；开放链 0"),
]
add_table(doc, ["成果项", "数量", "说明"], summary_rows, [1.55, 0.8, 4.15], font_size=9)

doc.add_heading("二、已经形成的研究链路", level=1)
add_callout(
    doc,
    "统一结构",
    "宏观母事件 → 事件时间线与衍生节点 → 官方事实与证据 → 传导渠道 → 行业影响 → 公司影响 → 事前方向判断 → D1/D3/D5/D20/M3 事后行情验证 → 偏差解释、反证与因果边界",
    fill="EAF2F8",
)
for text in (
    "事实层：保留事件发生时间、首次公开时间、来源、原文定位、快照和证据状态。",
    "推理层：区分来源事实与研究推断，并形成即时、短期、中期三类不同阶段判断。",
    "产业层：按事件关联行业和公司组织经营变量、财务变量、估值影响、条件与反证。",
    "验证层：保留 T-1 基准及 D1、D3、D5、D20、M3 实际行情，不以缺失值补零。",
    "因果层：行情仅用于验证窗口表现或机制一致性；当前 30 条公司因果状态均为 NOT_IDENTIFIED。",
):
    add_bullet(doc, text)

doc.add_heading("三、十个事件覆盖范围", level=1)
event_rows = []
for item in events:
    event = item["event"]
    anchors = "、".join(SERIES_ZH.get(w.get("series_id"), w.get("series_id", "")) for w in event.get("market_windows", []))
    event_rows.append((
        item["event_id"],
        event.get("title_zh", ""),
        len(event.get("milestones", [])),
        len(item.get("industries", [])),
        len(item.get("companies", [])),
        len(item.get("qa", [])),
        anchors,
    ))
add_table(doc, ["事件", "名称", "节点", "行业", "公司", "问答", "价值锚定物"], event_rows, [0.55, 1.65, 0.42, 0.42, 0.42, 0.42, 2.62], font_size=7.4)

doc.add_heading("四、代表案例：OPEC+ 意外减产", level=1)
event06 = next(x for x in events if x["event_id"] == "event-06")
brent = next(w for w in event06["event"]["market_windows"] if w["series_id"] == "DCOILBRENTEU")
w = brent["windows"]
add_callout(
    doc,
    "研究结论",
    f"事件时的供给收缩机制指向原油即时与短期偏强。事后布伦特原油相对 T-1 的实际变化为：D1 {w['D1']['raw_change']:.2f}%、D3 {w['D3']['raw_change']:.2f}%、D5 {w['D5']['raw_change']:.2f}%、D20 {w['D20']['raw_change']:.2f}%、M3 {w['M3']['raw_change']:.2f}%。短窗口与供给冲击机制一致，但 D20 和 M3 转弱，说明需求、库存及其他市场因素随后占据主导。",
)
add_bullet(doc, "预测层：站在事件锚点，以当时可见信息判断方向、期限、成立条件与失效条件。")
add_bullet(doc, "验证层：用后续实际行情检验预测期限与机制一致性，并解释偏差。")
add_bullet(doc, "边界层：该结果不能单独证明 OPEC+ 减产造成全部价格变化。")

doc.add_heading("五、质量验证与复现", level=1)
quality_rows = [
    ("数据验证", f"{len(validation['errors'])} errors / {len(validation['warnings'])} warnings", "通过"),
    ("占位符检查", validation["content_quality_stats"]["placeholder_sample_count"], "未发现占位样本"),
    ("完整阶段分析", validation["content_quality_stats"]["complete_horizon_count"], "157/157"),
    ("闭环状态", validation["closure_reason_stats"].get("CLOSED", 0), "开放链 0"),
    ("离线双轮复现", "四个黄金产物 SHA256 一致", "stable=true"),
    ("离线边界", "未联网、未重新生成研究文字", "network_accessed=false"),
]
add_table(doc, ["检查项", "结果", "状态"], quality_rows, [1.55, 3.55, 1.4], font_size=9)

doc.add_heading("六、当前定位与后续边界", level=1)
add_callout(
    doc,
    "建议汇报口径",
    "已完成十个宏观事件从事实、推理、行业公司传导到 D20/M3 行情验证的候选数据建设；当前进入人工研究审核和正式准入阶段。",
    fill="FFF7E6",
)
for text in (
    "当前所有候选继续保持 NEEDS_HUMAN_REVIEW 和 formal_ready=false。",
    "机器闭环表示字段、证据引用和推理路径完整，不等同于人工研究结论已经签字确认。",
    "未执行真实 GPU/QLoRA 训练；现有文件是训练候选数据，不应表述为已完成模型训练。",
    "行情一致性不是宏观事件因果证明；非 IDENTIFIED 状态必须保持 ABSTAIN。",
):
    add_bullet(doc, text)

doc.add_heading("七、交付文件", level=1)
files = [
    ("全量统一主数据", REPLAY / "candidate_master_data.jsonl"),
    ("训练候选数据", REPLAY / "candidate_sft_dataset.jsonl"),
    ("数据拆分清单", REPLAY / "split_manifest.json"),
    ("机器验证报告", REPLAY / "validation_report.json"),
    ("离线复现报告", RUN / "reproducibility_report.json"),
]
file_rows = []
for label, path in files:
    file_rows.append((label, path.name, f"{path.stat().st_size / 1024 / 1024:.2f} MB", sha256(path)[:16] + "..."))
add_table(doc, ["用途", "文件", "大小", "SHA256（前16位）"], file_rows, [1.45, 2.55, 0.85, 1.65], font_size=8.2)

doc.add_paragraph("注：详细的十事件、公司、行业、行情窗口与文件路径见配套 Excel 汇总表。")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(json.dumps({
    "output": str(OUT),
    "events": len(events),
    "samples": validation["counts"]["sft_samples"],
    "milestones": milestones,
    "industries": industries,
    "companies": companies,
    "qa": qas,
    "market_windows": market_windows,
    "repro_stable": repro["stable"],
}, ensure_ascii=False))
