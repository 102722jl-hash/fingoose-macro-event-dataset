"""Source parsers and immutable snapshot helpers for commodity market v0.2."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import urllib.request
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Mapping, Sequence

import xlrd
from openpyxl import load_workbook

from fingoose_commodity_market import (
    COMMODITY_SPECS,
    ContractBar,
    MarketDataError,
    TradingCalendar,
)


ADAPTER_VERSION = "fingoose-official-exchange-adapter-v0.2"

SOURCE_CONFIG = {
    "SHFE": {
        "base_url": "https://www.shfe.com.cn",
        "provider": "Shanghai Futures Exchange",
        "landing_url": "https://www.shfe.com.cn/reports/tradedata/dailyandweeklydata/",
        "commodities": ("CU", "AU"),
    },
    "INE": {
        "base_url": "https://www.ine.cn",
        "provider": "Shanghai International Energy Exchange",
        "landing_url": "https://www.ine.cn/reports/tradedata/dailyandweeklydata/",
        "commodities": ("SC",),
    },
    "DCE": {
        "provider": "Dalian Commodity Exchange",
        "landing_url": "https://www.dce.com.cn/dalianshangpin/xqsj/tjsj26/rtj/rttj/index.html",
        "commodities": ("I",),
    },
}


class AdapterError(ValueError):
    """Raised when source data cannot be normalized without guessing."""


@dataclass(frozen=True)
class ParsedContract:
    row: dict[str, Any]
    quality_flags: tuple[str, ...] = ()


@dataclass
class ParseResult:
    records: list[ParsedContract]
    rejected: list[dict[str, Any]]


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def json_bytes(value: Any) -> bytes:
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def stable_record_id(payload: Mapping[str, Any]) -> str:
    return "FGCMD-" + sha256(json_bytes(payload))[:20].upper()


def safe_name(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", value)
    return cleaned.strip("._") or "snapshot.bin"


class SnapshotStore:
    """Content-addressed immutable raw snapshot storage."""

    def __init__(self, root: Path):
        self.root = root

    def store(
        self,
        *,
        source_id: str,
        provider: str,
        source_url: str,
        requested_date: str | None,
        content: bytes,
        original_name: str,
        content_type: str,
        http_status: int | None,
        retrieved_at: str | None = None,
        adapter_version: str = ADAPTER_VERSION,
    ) -> dict[str, Any]:
        retrieved_at = retrieved_at or now_iso()
        digest = sha256(content)
        directory = self.root / safe_name(source_id) / (requested_date or "undated")
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / f"{digest[:16]}_{safe_name(original_name)}"
        if path.exists():
            if path.read_bytes() != content:
                raise AdapterError(f"snapshot hash collision at {path}")
        else:
            with path.open("xb") as handle:
                handle.write(content)
        identity = {
            "source_id": source_id,
            "source_url": source_url,
            "requested_date": requested_date,
            "content_sha256": digest,
            "adapter_version": adapter_version,
        }
        return {
            "manifest_version": "fingoose-commodity-source-manifest-v0.2",
            "manifest_id": "FGSRC-" + sha256(json_bytes(identity))[:20].upper(),
            "snapshot_id": "FGSNAP-" + digest[:20].upper(),
            "source_id": source_id,
            "provider": provider,
            "source_url": source_url,
            "requested_date": requested_date,
            "retrieved_at": retrieved_at,
            "original_name": original_name,
            "local_path": path.as_posix(),
            "http_status": http_status,
            "content_type": content_type,
            "byte_count": len(content),
            "content_sha256": digest,
            "adapter_version": adapter_version,
            "license_scope": "unknown",
            "display_allowed": False,
            "redistribution_allowed": False,
            "parse_status": "collected",
            "record_count": 0,
            "errors": [],
            "revision": 1,
            "is_latest": True,
        }

    def store_file(
        self,
        *,
        source_id: str,
        provider: str,
        source_url: str,
        requested_date: str | None,
        source_path: Path,
        content_type: str,
        http_status: int | None,
        retrieved_at: str | None = None,
        adapter_version: str = ADAPTER_VERSION,
        chunk_size: int = 1024 * 1024,
    ) -> dict[str, Any]:
        """Stream a local source into immutable content-addressed storage."""

        if chunk_size < 1:
            raise AdapterError("chunk_size must be positive")
        if not source_path.is_file():
            raise AdapterError(f"snapshot source file is missing: {source_path}")

        def file_digest(path: Path) -> tuple[str, int]:
            digest = hashlib.sha256()
            byte_count = 0
            with path.open("rb") as handle:
                while chunk := handle.read(chunk_size):
                    digest.update(chunk)
                    byte_count += len(chunk)
            return digest.hexdigest(), byte_count

        digest, byte_count = file_digest(source_path)
        directory = self.root / safe_name(source_id) / (requested_date or "undated")
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / f"{digest[:16]}_{safe_name(source_path.name)}"
        if path.exists():
            existing_digest, existing_size = file_digest(path)
            if existing_digest != digest or existing_size != byte_count:
                raise AdapterError(f"snapshot hash collision at {path}")
        else:
            try:
                with source_path.open("rb") as source, path.open("xb") as target:
                    while chunk := source.read(chunk_size):
                        target.write(chunk)
            except Exception:
                if path.exists():
                    path.unlink()
                raise

        retrieved_at = retrieved_at or now_iso()
        identity = {
            "source_id": source_id,
            "source_url": source_url,
            "requested_date": requested_date,
            "content_sha256": digest,
            "adapter_version": adapter_version,
        }
        return {
            "manifest_version": "fingoose-commodity-source-manifest-v0.2",
            "manifest_id": "FGSRC-" + sha256(json_bytes(identity))[:20].upper(),
            "snapshot_id": "FGSNAP-" + digest[:20].upper(),
            "source_id": source_id,
            "provider": provider,
            "source_url": source_url,
            "requested_date": requested_date,
            "retrieved_at": retrieved_at,
            "original_name": source_path.name,
            "local_path": path.as_posix(),
            "http_status": http_status,
            "content_type": content_type,
            "byte_count": byte_count,
            "content_sha256": digest,
            "adapter_version": adapter_version,
            "license_scope": "unknown",
            "display_allowed": False,
            "redistribution_allowed": False,
            "parse_status": "collected",
            "record_count": 0,
            "errors": [],
            "revision": 1,
            "is_latest": True,
        }


def download(url: str, timeout: int = 30) -> tuple[bytes, str, int]:
    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 FinGoose/0.2",
            "Referer": url.split("/data/", 1)[0] + "/",
        },
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        content = response.read()
        content_type = response.headers.get("Content-Type", "application/octet-stream")
        status = int(getattr(response, "status", 200))
    if not content:
        raise AdapterError(f"empty response from {url}")
    return content, content_type, status


def parse_json_bytes(content: bytes, source_name: str) -> dict[str, Any]:
    try:
        value = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise AdapterError(f"invalid JSON from {source_name}") from exc
    if not isinstance(value, dict):
        raise AdapterError(f"expected an object from {source_name}")
    return value


def yyyymmdd(value: Any, field_name: str) -> str:
    text = str(value).strip()
    if not re.fullmatch(r"\d{8}", text):
        raise AdapterError(f"invalid {field_name}: {value!r}")
    return date(int(text[:4]), int(text[4:6]), int(text[6:8])).isoformat()


def positive_number(value: Any, field_name: str) -> float:
    text = "" if value is None else str(value).strip().replace(",", "")
    if text in {"", "-", "--", "nan", "None"}:
        raise AdapterError(f"missing {field_name}")
    try:
        number = float(text)
    except ValueError as exc:
        raise AdapterError(f"invalid {field_name}: {value!r}") from exc
    if not math.isfinite(number) or number <= 0:
        raise AdapterError(f"invalid {field_name}: {value!r}")
    return number


def nonnegative_integer(value: Any, field_name: str) -> int:
    text = "" if value is None else str(value).strip().replace(",", "")
    if text in {"", "-", "--", "nan", "None"}:
        raise AdapterError(f"missing {field_name}")
    try:
        number = float(text)
    except ValueError as exc:
        raise AdapterError(f"invalid {field_name}: {value!r}") from exc
    if not math.isfinite(number) or number < 0 or not number.is_integer():
        raise AdapterError(f"invalid {field_name}: {value!r}")
    return int(number)



def parse_shfe_like(
    daily_payload: Mapping[str, Any],
    contract_payload: Mapping[str, Any],
    *,
    exchange: str,
    commodity_code: str,
    trade_date: str,
) -> ParseResult:
    if commodity_code not in SOURCE_CONFIG[exchange]["commodities"]:
        raise AdapterError(f"{commodity_code} is not configured for {exchange}")
    prefix = COMMODITY_SPECS[commodity_code]["contract_prefix"]
    metadata_rows = contract_payload.get("ContractBaseInfo")
    daily_rows = daily_payload.get("o_curinstrument")
    if not isinstance(metadata_rows, list) or not isinstance(daily_rows, list):
        raise AdapterError("source payload is missing expected arrays")
    metadata = {}
    for item in metadata_rows:
        instrument = str(item.get("INSTRUMENTID", "")).strip().lower()
        if re.fullmatch(re.escape(prefix) + r"\d{4}", instrument):
            metadata[instrument] = item

    records: list[ParsedContract] = []
    rejected: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, item in enumerate(daily_rows):
        group = str(item.get("PRODUCTGROUPID", "")).strip().lower()
        product_id = str(item.get("PRODUCTID", "")).strip().lower()
        delivery = str(item.get("DELIVERYMONTH", "")).strip()
        if group != prefix:
            continue
        if (
            str(item.get("PRODUCTCLASS", "1")).strip() != "1"
            or "tas" in product_id
            or not re.fullmatch(r"\d{4}", delivery)
        ):
            rejected.append({"row_index": index, "reason": "non_contract_or_aggregate"})
            continue
        contract_code = prefix + delivery
        if contract_code in seen:
            raise AdapterError(f"duplicate source row: {trade_date} {contract_code}")
        seen.add(contract_code)
        contract_meta = metadata.get(contract_code)
        if contract_meta is None:
            raise AdapterError(f"missing ContractBaseInfo for {contract_code}")
        metadata_day = str(contract_meta.get("TRADINGDAY", "")).strip()
        if metadata_day and metadata_day != trade_date.replace("-", ""):
            raise AdapterError(f"metadata trading day mismatch for {contract_code}")
        flags: list[str] = []
        raw_exchange = str(contract_meta.get("EXCHANGEID", "")).strip()
        if raw_exchange and raw_exchange != exchange:
            flags.append("SOURCE_EXCHANGE_ID_MISMATCH")
        try:
            row = {
                "commodity_code": commodity_code,
                "exchange": exchange,
                "trade_date": trade_date,
                "contract_code": contract_code,
                "contract_month": f"20{delivery[:2]}-{delivery[2:]}",
                "last_trade_date": yyyymmdd(
                    contract_meta.get("EXPIREDATE"), "EXPIREDATE"
                ),
                "open": positive_number(item.get("OPENPRICE"), "OPENPRICE"),
                "high": positive_number(item.get("HIGHESTPRICE"), "HIGHESTPRICE"),
                "low": positive_number(item.get("LOWESTPRICE"), "LOWESTPRICE"),
                "close": positive_number(item.get("CLOSEPRICE"), "CLOSEPRICE"),
                "settlement": positive_number(
                    item.get("SETTLEMENTPRICE"), "SETTLEMENTPRICE"
                ),
                "pre_settlement": positive_number(
                    item.get("PRESETTLEMENTPRICE"), "PRESETTLEMENTPRICE"
                ),
                "volume": nonnegative_integer(item.get("VOLUME"), "VOLUME"),
                "open_interest": nonnegative_integer(
                    item.get("OPENINTEREST"), "OPENINTEREST"
                ),
                "is_trading": True,
            }
            ContractBar.from_mapping(row)
        except (AdapterError, MarketDataError, ValueError) as exc:
            rejected.append(
                {"row_index": index, "contract_code": contract_code, "reason": str(exc)}
            )
            continue
        records.append(ParsedContract(row, tuple(flags)))
    if not records:
        raise AdapterError(f"no valid {commodity_code} contracts on {trade_date}")
    return ParseResult(records, rejected)


def parse_trade_data_js(
    content: bytes, *, exchange: str, year: int, coverage_end: str
) -> TradingCalendar:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise AdapterError("trade-data.js is not UTF-8") from exc
    key = rf"(?:^|[{{,])\s*['\"]?{year}['\"]?\s*:\s*"
    quoted = re.search(key + r"(?P<quote>['\"])(?P<body>.*?)(?P=quote)", text)
    array = re.search(key + r"\[(?P<body>.*?)\]", text, flags=re.DOTALL)
    match = quoted or array
    if not match:
        raise AdapterError(f"trade-data.js has no {year} entry")
    tokens = re.findall(r"(?<!\d)(?:\d{8}|\d{4})(?!\d)", match.group("body"))
    listed: set[date] = set()
    for token in tokens:
        compact = token if len(token) == 8 else f"{year}{token}"
        try:
            parsed = date(
                int(compact[:4]), int(compact[4:6]), int(compact[6:8])
            )
        except ValueError:
            continue
        if parsed.year == year:
            listed.add(parsed)
    if not listed:
        raise AdapterError(f"trade-data.js {year} entry contains no dates")

    requested_end = date.fromisoformat(coverage_end)
    if requested_end.year != year:
        raise AdapterError("trade-data.js parser accepts one calendar year at a time")
    start = date(year, 1, 1)
    listed_end = min(max(listed), requested_end)
    weekdays_to_listed_end = {
        start + timedelta(days=offset)
        for offset in range((listed_end - start).days + 1)
        if (start + timedelta(days=offset)).weekday() < 5
    }
    listed_weekdays = {item for item in listed if item.weekday() < 5}
    density = len(listed_weekdays & weekdays_to_listed_end) / max(
        len(weekdays_to_listed_end), 1
    )
    lists_trading_days = density >= 0.6
    actual_end = listed_end if lists_trading_days else requested_end
    weekdays = {
        start + timedelta(days=offset)
        for offset in range((actual_end - start).days + 1)
        if (start + timedelta(days=offset)).weekday() < 5
    }
    available = weekdays & listed if lists_trading_days else weekdays - listed
    return TradingCalendar.from_values(
        exchange,
        sorted(available),
        coverage_start=start,
        coverage_end=actual_end,
        verified=False,
    )


def normalized_header(value: Any) -> str:
    return re.sub(r"[\s\u3000]+", "", str(value)).lower()


DCE_ALIASES = {
    "product_name": {"商品名称", "品种", "品种名称", "productname"},
    "delivery_month": {"交割月份", "合约", "合约代码", "deliverymonth", "contract"},
    "open": {"开盘价", "开盘", "open", "openprice"},
    "high": {"最高价", "最高", "high", "highestprice"},
    "low": {"最低价", "最低", "low", "lowestprice"},
    "close": {"收盘价", "收盘", "close", "closeprice"},
    "pre_settlement": {
        "前结算价", "昨结算", "presettlement", "presettlementprice"
    },
    "settlement": {"结算价", "结算", "settlement", "settlementprice"},
    "volume": {"成交量", "成交手", "volume"},
    "open_interest": {"持仓量", "持仓手", "openinterest"},
}


def detect_table_format(content: bytes) -> str:
    prefix = content[:32].lstrip()
    if content.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        return "xls"
    if content.startswith(b"PK\x03\x04"):
        return "xlsx"
    if prefix.lower().startswith((b"<html", b"<!doctype", b"<table")):
        return "html"
    return "csv"


def detected_table_content_type(content: bytes) -> str:
    return {
        "csv": "text/csv",
        "html": "text/html",
        "xls": "application/vnd.ms-excel",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }[detect_table_format(content)]


class _FirstHtmlTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.in_table = False
        self.in_cell = False
        self.finished = False
        self.current_cell: list[str] = []
        self.current_row: list[str] = []
        self.rows: list[list[str]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag == "table" and not self.in_table and not self.finished:
            self.in_table = True
        elif self.in_table and tag == "tr":
            self.current_row = []
        elif self.in_table and tag in {"td", "th"}:
            self.in_cell = True
            self.current_cell = []

    def handle_data(self, data: str) -> None:
        if self.in_cell:
            self.current_cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if self.in_table and tag in {"td", "th"}:
            self.current_row.append("".join(self.current_cell).strip())
            self.in_cell = False
        elif self.in_table and tag == "tr" and self.current_row:
            self.rows.append(self.current_row)
            self.current_row = []
        elif self.in_table and tag == "table":
            self.in_table = False
            self.finished = True


def _rows_to_dicts(rows: Sequence[Sequence[Any]], source_name: str) -> list[dict[str, Any]]:
    if not rows:
        raise AdapterError(f"empty table in {source_name}")
    headers = [str(value).strip() for value in rows[0]]
    if not headers or any(not value for value in headers):
        raise AdapterError(f"blank table header in {source_name}")
    if len(set(headers)) != len(headers):
        raise AdapterError(f"duplicate table header in {source_name}")
    result = []
    for values in rows[1:]:
        if not any(value not in {None, ""} for value in values):
            continue
        padded = list(values) + [None] * (len(headers) - len(values))
        result.append(dict(zip(headers, padded[: len(headers)])))
    return result


def read_table_bytes(content: bytes, original_name: str) -> list[dict[str, Any]]:
    detected = detect_table_format(content)
    if detected == "xls":
        workbook = xlrd.open_workbook(file_contents=content)
        sheet = workbook.sheet_by_index(0)
        rows = [sheet.row_values(index) for index in range(sheet.nrows)]
        return _rows_to_dicts(rows, original_name)
    if detected == "xlsx":
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True
        )
        sheet = workbook.active
        rows = [list(values) for values in sheet.iter_rows(values_only=True)]
        workbook.close()
        return _rows_to_dicts(rows, original_name)
    if detected == "html":
        text = None
        for encoding in ("utf-8", "gb18030"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise AdapterError(f"cannot decode HTML table {original_name}")
        parser = _FirstHtmlTableParser()
        parser.feed(text)
        return _rows_to_dicts(parser.rows, original_name)
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            continue
        reader = csv.DictReader(io.StringIO(text))
        if reader.fieldnames:
            return [dict(row) for row in reader]
    raise AdapterError(f"cannot parse DCE table {original_name}")


def column_map(rows: Sequence[Mapping[str, Any]]) -> dict[str, str]:
    if not rows:
        raise AdapterError("DCE table is empty")
    normalized = {normalized_header(column): column for column in rows[0].keys()}
    result: dict[str, str] = {}
    for canonical, aliases in DCE_ALIASES.items():
        accepted = {normalized_header(alias) for alias in aliases}
        matches = [
            original for name, original in normalized.items() if name in accepted
        ]
        if len(matches) != 1:
            raise AdapterError(f"DCE table needs exactly one {canonical} column")
        result[canonical] = str(matches[0])
    return result

def load_dce_metadata(path: Path) -> dict[str, dict[str, str]]:
    if not path.exists():
        raise AdapterError(f"DCE metadata is missing: {path}")
    content = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise AdapterError("DCE metadata encoding is unsupported")
    rows = csv.DictReader(io.StringIO(text))
    required = {"contract_code", "last_trade_date", "source_url", "retrieved_at"}
    if rows.fieldnames is None or not required.issubset(rows.fieldnames):
        raise AdapterError(f"DCE metadata requires columns: {sorted(required)}")
    result = {}
    for row in rows:
        code = str(row["contract_code"]).strip().lower()
        if not re.fullmatch(r"i\d{4}", code):
            raise AdapterError(f"invalid DCE contract_code: {code!r}")
        try:
            last_trade = date.fromisoformat(row["last_trade_date"])
            retrieved = datetime.fromisoformat(
                row["retrieved_at"].replace("Z", "+00:00")
            )
        except ValueError as exc:
            raise AdapterError(f"invalid DCE metadata dates for {code}") from exc
        if retrieved.tzinfo is None:
            raise AdapterError(f"DCE retrieved_at lacks timezone for {code}")
        if not row["source_url"].startswith("https://"):
            raise AdapterError(f"DCE source_url must use HTTPS for {code}")
        if code in result:
            raise AdapterError(f"duplicate DCE metadata for {code}")
        result[code] = {**row, "last_trade_date": last_trade.isoformat()}
    return result


def parse_dce_table(
    frame: Sequence[Mapping[str, Any]],
    *,
    trade_date: str,
    metadata: Mapping[str, Mapping[str, str]],
) -> ParseResult:
    columns = column_map(frame)
    records: list[ParsedContract] = []
    rejected: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, item in enumerate(frame):
        product_name = str(item[columns["product_name"]]).strip()
        raw_contract = str(item[columns["delivery_month"]]).strip().lower()
        if product_name in {"nan", "", "总计", "小计"}:
            continue
        if "铁矿石" not in product_name and not raw_contract.startswith("i"):
            continue
        raw_contract = raw_contract.removesuffix(".0")
        if re.fullmatch(r"\d{4}", raw_contract):
            contract_code = "i" + raw_contract
        elif re.fullmatch(r"i\d{4}", raw_contract):
            contract_code = raw_contract
        else:
            rejected.append({"row_index": index, "reason": "invalid_contract_code"})
            continue
        if contract_code in seen:
            raise AdapterError(f"duplicate DCE row: {trade_date} {contract_code}")
        seen.add(contract_code)
        contract_meta = metadata.get(contract_code)
        if contract_meta is None:
            raise AdapterError(f"missing DCE metadata for {contract_code}")
        suffix = contract_code[-4:]
        try:
            row = {
                "commodity_code": "I",
                "exchange": "DCE",
                "trade_date": trade_date,
                "contract_code": contract_code,
                "contract_month": f"20{suffix[:2]}-{suffix[2:]}",
                "last_trade_date": contract_meta["last_trade_date"],
                "open": positive_number(item[columns["open"]], "open"),
                "high": positive_number(item[columns["high"]], "high"),
                "low": positive_number(item[columns["low"]], "low"),
                "close": positive_number(item[columns["close"]], "close"),
                "settlement": positive_number(
                    item[columns["settlement"]], "settlement"
                ),
                "pre_settlement": positive_number(
                    item[columns["pre_settlement"]], "pre_settlement"
                ),
                "volume": nonnegative_integer(item[columns["volume"]], "volume"),
                "open_interest": nonnegative_integer(
                    item[columns["open_interest"]], "open_interest"
                ),
                "is_trading": True,
            }
            ContractBar.from_mapping(row)
        except (AdapterError, MarketDataError, ValueError) as exc:
            rejected.append(
                {"row_index": index, "contract_code": contract_code, "reason": str(exc)}
            )
            continue
        records.append(ParsedContract(row))
    if not records:
        raise AdapterError(f"no valid iron-ore contracts on {trade_date}")
    return ParseResult(records, rejected)


def trade_date_from_filename(path: Path) -> str:
    matches = re.findall(r"(?<!\d)(20\d{6})(?!\d)", path.stem)
    if len(matches) != 1:
        raise AdapterError(
            f"DCE filename must contain one YYYYMMDD date: {path.name}"
        )
    return yyyymmdd(matches[0], "DCE filename date")


def bundle_provenance(
    manifests: Sequence[Mapping[str, Any]], *, commodity_code: str
) -> dict[str, Any]:
    if not manifests:
        raise AdapterError(f"no source manifests for {commodity_code}")
    exchange = COMMODITY_SPECS[commodity_code]["exchange"]
    identity = sorted(
        (
            {
                "source_id": item["source_id"],
                "requested_date": item["requested_date"],
                "content_sha256": item["content_sha256"],
            }
            for item in manifests
        ),
        key=lambda item: (
            str(item["source_id"]),
            str(item["requested_date"]),
            str(item["content_sha256"]),
        ),
    )
    digest = sha256(json_bytes(identity))
    return {
        "source_id": (
            "dce_manual_bundle"
            if exchange == "DCE"
            else f"{exchange.lower()}_official_bundle"
        ),
        "source_url": SOURCE_CONFIG[exchange]["landing_url"],
        "provider": SOURCE_CONFIG[exchange]["provider"],
        "retrieved_at": max(str(item["retrieved_at"]) for item in manifests),
        "snapshot_id": f"FGBUNDLE-{commodity_code}-{digest[:16].upper()}",
        "content_sha256": digest,
        "adapter_version": ADAPTER_VERSION,
        "license_scope": "unknown",
        "display_allowed": False,
        "redistribution_allowed": False,
    }


def manifest_signature(record: Mapping[str, Any]) -> tuple[str, str]:
    return str(record["content_sha256"]), str(record["adapter_version"])


def merge_source_manifest_revisions(
    existing: Sequence[Mapping[str, Any]],
    candidates: Sequence[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    merged = []
    for value in existing:
        record = dict(value)
        record.setdefault("revision", 1)
        record.setdefault("is_latest", True)
        merged.append(record)
    by_key: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
    for record in merged:
        key = (
            str(record["source_id"]),
            str(record["source_url"]),
            str(record["requested_date"]),
        )
        by_key.setdefault(key, []).append(record)
    for candidate_value in candidates:
        candidate = dict(candidate_value)
        key = (
            str(candidate["source_id"]),
            str(candidate["source_url"]),
            str(candidate["requested_date"]),
        )
        history = by_key.setdefault(key, [])
        latest_records = [item for item in history if item.get("is_latest")]
        if len(latest_records) > 1:
            raise AdapterError(f"multiple latest source manifests for {key}")
        latest = latest_records[0] if latest_records else None
        if latest is not None and manifest_signature(latest) == manifest_signature(candidate):
            continue
        if latest is not None:
            latest["is_latest"] = False
        candidate["revision"] = max(
            (int(item.get("revision", 0)) for item in history), default=0
        ) + 1
        candidate["is_latest"] = True
        merged.append(candidate)
        history.append(candidate)
    return sorted(
        merged,
        key=lambda item: (
            str(item["source_id"]),
            str(item["requested_date"]),
            int(item["revision"]),
            str(item["content_sha256"]),
        ),
    )

def record_signature(record: Mapping[str, Any]) -> str:
    payload = {
        key: value
        for key, value in record.items()
        if key not in {"record_id", "revision", "is_latest", "quality"}
    }
    provenance = dict(payload.get("provenance", {}))
    provenance.pop("retrieved_at", None)
    payload["provenance"] = provenance
    return sha256(json_bytes(payload))


def recompute_record_id(record: dict[str, Any]) -> None:
    payload = {
        key: value for key, value in record.items() if key not in {"record_id", "quality"}
    }
    record["record_id"] = stable_record_id(payload)


def merge_contract_revisions(
    existing: Sequence[Mapping[str, Any]],
    candidates: Sequence[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    merged = [dict(item) for item in existing]
    by_key: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
    for record in merged:
        key = (
            record["commodity"]["exchange"],
            record["contract_code"],
            record["trade_date"],
        )
        by_key.setdefault(key, []).append(record)
    for candidate_value in candidates:
        candidate = dict(candidate_value)
        key = (
            candidate["commodity"]["exchange"],
            candidate["contract_code"],
            candidate["trade_date"],
        )
        history = by_key.setdefault(key, [])
        latest = next((item for item in history if item.get("is_latest")), None)
        if latest is not None and record_signature(latest) == record_signature(candidate):
            continue
        if latest is not None:
            latest["is_latest"] = False
            recompute_record_id(latest)
        candidate["revision"] = max(
            (int(item.get("revision", 0)) for item in history), default=0
        ) + 1
        candidate["is_latest"] = True
        recompute_record_id(candidate)
        merged.append(candidate)
        history.append(candidate)
    return sorted(
        merged,
        key=lambda item: (
            item["commodity"]["commodity_code"],
            item["trade_date"],
            item["contract_code"],
            item["revision"],
        ),
    )
